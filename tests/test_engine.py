# -*- coding: utf-8 -*-
"""가상 원시자료(실제 이상사례 아님)로 A단계 → 사용자 입력 → B단계 전체 흐름을 검증"""
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import engine  # noqa: E402

CODE = "200000001"
TXT = {
    "README": f"<KAERS DB 제공파일 정보>\n\n요청 품목코드 : {CODE}\n보고건수 : 2\n",
    "DEMO": (
        "KAERS_NO|DEPT_RECEIPT_NO|SFRPNO|RPT_DL_DT|KAERS_GB|RPT_TY|ADRSE_STUDY_TYP|ADRSE_STUDY_LWPRT_TYP|LTRTRE_INFO|"
        "CLIENT_STUDY_NO|FIRST_OCCR_DT|RECENT_OCCR_DT|QCK_RPT_YN|SFRPNO_2|REPRT_CHANGE_CD|PRMRPT_TY|PRMRPT_LWPRT_CD|"
        "SENDER_TY|SENDER_TY_MED_EXPERT|PTNT_OCCURSYMT_AGE|PTNT_OCCURSYMT_AGE_UNIT|PTNT_AGRDE|PTNT_BRTYR_YYYY|PTNT_SEX|"
        "PTNT_WEIGHT|PTNT_HEIGHT|OCCURSYMT_PREG_TRM|OCCURSYMT_PREG_TRM_UNIT\n"
        "1001|D1|S1|20210318|1|1||||||||||1||4||45|00103|5|1976|1|81.1|||\n"
        "1002|D2|S2|20220101|1|2|3|1|문헌A|P-01|||Y||2|UNK||3|1|7|00107|1||2||||\n"
    ),
    "HIST_E": "KAERS_NO|PTNT_HIST_MEDDRA_KOR_NM\n1001|\n",
    "PARENT": "KAERS_NO|PARENT_AGE\n1001|\n",
    "EVENT": (
        "KAERS_NO|ADR_SEQ|ADR_MEDDRA_KOR_NM|ADR_MEDDRA_ENG_NM|ADR_START_DT|ADR_END_DT|ADR_RESULT_CODE|"
        "CLNIC_FACT_CONFIRM_YN|WHOART_ARRN|WHOART_SEQ|SE_DEATH|SE_LIFE_MENACE|SE_HSPTLZ_EXTN|SE_FNCT_DGRD|SE_ANMLY|SE_ETC_IMPRTNC_SITTN\n"
        "1001|1|소양증|Pruritus|20200617|20200618|1|Y|0024|001||||||\n"
        "1002|1|두통|Headache|20220101||3|Y|0001|002||Y|||||\n"
    ),
    "TEST": "KAERS_NO|TEST_DATE\n1001|\n",
    "DRUG": (
        "KAERS_NO|DRUG_SEQ|DRUG_GB|DRUG_CD|ACCMLT_DOSAGE_QTY|ACCMLT_DOSAGE_QTY_UNIT|DRUG_ACTION\n"
        f"1001|1|1|{CODE}|||1\n"
        "1001|2|2|299999999|||4\n"
        f"1002|1|2|{CODE}|||9\n"
    ),
    "DRUG1": "KAERS_NO|DRUG_SEQ|INGR_SEQ|INGR_CD\n1001|1|1|X\n",
    "DRUG2": "KAERS_NO|DRUG_SEQ|DOSAGE_QTY\n1001|1|1\n",
    "DRUG3": "KAERS_NO|DRUG_SEQ|DSAS_CD\n1001|1|\n",
    "DRUG_EVENT": "KAERS_NO|DRUG_SEQ|ADR_SEQ|RE_DOSAGE_ALLERGY_YN\n1001|1|1|4\n1001|2|1|\n1002|1|1|\n",
    "ASSESSMENT": "KAERS_NO|DRUG_SEQ|ADR_SEQ|EVALT_RESULT_CODE\n1001|1|1|3\n1002|1|1|\n",
    "GROUP": "KAERS_NO|GROUP|SEQ\n1001|9001|1\n1002|9002|1\n",
}


def sheet_dicts(ws):
    rows = list(ws.iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:]]


class EngineFlowTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        root = Path(cls.tmp.name)
        txt = root / "files"; txt.mkdir()
        for name, body in TXT.items():
            (txt / f"{name}.txt").write_text(body, encoding="utf-8-sig")
        cls.out = root
        cls.path03 = engine.run_step_a(txt, root, "테스트정2/500", "테스트정밀리그램", "20261015")
        # 사용자 입력 흉내: _03 EVENT 시트에 PT/SOC 입력
        wb = openpyxl.load_workbook(cls.path03)
        ev = wb["EVENT"]; h = [c.value for c in ev[1]]
        for r in range(2, ev.max_row + 1):
            ev.cell(r, h.index("PT_한국어") + 1, f"PT-{ev.cell(r, 1).value}")
            ev.cell(r, h.index("SOC_한국어") + 1, f"SOC-{ev.cell(r, 1).value}")
        wb.save(cls.path03)
        names = engine.drug_names_in(cls.path03)
        cls.names = names
        cls.path30 = engine.run_step_b(cls.path03, root, "테스트정2/500", names[0], "테스트정", "20261015")

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_file_names_safe(self):
        self.assertEqual(self.path03.name, "원시자료_테스트정2_500_03_20261015.xlsx")
        for no in ["01", "02", "03", "10", "20", "30"]:
            self.assertTrue((self.out / f"원시자료_테스트정2_500_{no}_20261015.xlsx").is_file(), no)
        self.assertEqual(engine.parse_out_name(self.path03)["product"], "테스트정2_500")

    def test_step3_codes(self):
        wb = openpyxl.load_workbook(self.path03)
        demo = sheet_dicts(wb["DEMO"])
        self.assertEqual(demo[0]["연령단위"], "년")          # '00103' 앞자리 0
        self.assertEqual(demo[1]["연령단위"], "일")
        self.assertEqual(demo[1]["원보고자"], "모름")        # UNK
        self.assertEqual(demo[1]["기타시험구분"], "재심사보고 사용성적조사")
        ev = sheet_dicts(wb["EVENT"])
        self.assertEqual([e["중대한"] for e in ev], [None, "Y"])
        self.assertEqual(ev[1]["최종관찰결과"], "회복되지 않았음")
        self.assertEqual(self.names, ["테스트정밀리그램"])

    def test_step4_lookup_and_user_inputs(self):
        wb = openpyxl.load_workbook(self.out / "원시자료_테스트정2_500_20_20261015.xlsx")
        rows = sheet_dicts(wb["Sheet1"])
        self.assertEqual(len(rows), 3)
        r = rows[0]
        self.assertEqual(r["부서접수번호"], "D1")
        self.assertEqual(r["의심약 구분"], "의심")
        self.assertEqual(r["의약품명"], "테스트정밀리그램")
        self.assertEqual(r["재투여 및 재발현 여부"], "재투여하지 않아 해당사항 없음")
        self.assertEqual(r["평가결과"], "가능함")
        self.assertEqual(r["PT_한국어"], "PT-1001")
        for row in rows:                                     # 사용자가 _30 에서 직접 입력: 제목만, 값은 공란
            self.assertIsNone(row["LLT_한국어"]); self.assertIsNone(row["반영여부"])
        self.assertEqual(r["최초추적그룹번호"], 9001)
        self.assertEqual(rows[1]["의약품명"], 0)             # 제품 외 병용약: 엑셀처럼 빈 셀 → 0
        self.assertEqual(rows[2]["중대한 이상사례(Y/0)"], "Y")
        self.assertEqual(wb["Sheet1"].freeze_panes, "D2")

    def test_split(self):
        wb = openpyxl.load_workbook(self.path30)
        self.assertEqual(wb.sheetnames, ["원본", "테스트정_의심", "테스트정_병용"])
        self.assertEqual(wb["테스트정_의심"].max_row - 1, 1)
        self.assertEqual(wb["테스트정_병용"].max_row - 1, 1)

    def test_missing_txt_error(self):
        with tempfile.TemporaryDirectory() as d:
            with self.assertRaises(engine.BotError):
                engine.run_step_a(Path(d), Path(d), "x", "x", "20261015")


if __name__ == "__main__":
    unittest.main()
