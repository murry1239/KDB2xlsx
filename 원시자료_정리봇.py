# -*- coding: utf-8 -*-
"""
원시자료 정리봇  (Python 3.11 / 필요 패키지: openpyxl)

KAERS 원시자료 txt 13개를 읽어 아래 6개 파일을 만듭니다.
  _01 : txt 13개 → 시트 13개
  _02 : 한글 설명 열 추가(빈 열)
  _03 : 코드값 → 한글 정보 입력
  _10 : DRUG_EVENT 시트에 다른 시트 정보 불러오기
  _20 : DRUG_EVENT 값만 새 파일로 (Sheet1)
  _30 : 원본 / 제품명_의심 / 제품명_병용 시트

사용법
  1) 명령 프롬프트에서:  python 원시자료_정리봇.py
  2) 화면 질문에 따라 작업 종류, txt 폴더, 제품명 등을 입력

작업 종류
  [1] 전체 실행      : txt → _01, _02, _03, _10, _20, _30
  [2] 1~3단계만      : txt → _01, _02, _03
                       (이후 _03 파일 EVENT 시트에 PT_한국어/SOC_한국어 직접 입력)
  [3] 4~6단계만      : PT/SOC 입력을 마친 _03 파일 → _10, _20, _30

명세 대비 수정 사항
  - 4단계 수식 중 부서접수번호/보고자관리번호/보고완료일의
    MATCH(A2,DEMO!A:A),0) 괄호 오류를 MATCH(A2,DEMO!A:A,0) (정확히 일치)로 바로잡아 적용
  - 4단계는 엑셀 수식 대신 파이썬에서 직접 값을 계산 (원래 수식은 DRUG_EVENT_LOOKUPS에 주석으로 보존)
  - 열 위치(P열, U열 등) 대신 열 이름으로 찾음 → 원본 열 구성이 바뀌어도 안전
"""
from __future__ import annotations

import csv
import datetime as dt
import os
import re
import sys
import traceback
from pathlib import Path

# 한글 Windows(cp949)에서 출력할 수 없는 문자가 있어도 멈추지 않도록 함
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(errors="replace")
    except (AttributeError, ValueError):
        pass

try:
    import openpyxl
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl 이 필요합니다.  명령 프롬프트에서  pip install openpyxl  실행 후 다시 시도하세요.")
    input("Enter 를 누르면 종료합니다.")
    sys.exit(1)


# ════════════════════════════════════════════════════════════════════
# 설정
# ════════════════════════════════════════════════════════════════════
VERSION = "alpha_0.2"

SHEET_ORDER = ["README", "DEMO", "HIST_E", "PARENT", "EVENT", "TEST", "DRUG",
               "DRUG1", "DRUG2", "DRUG3", "DRUG_EVENT", "ASSESSMENT", "GROUP"]

# 엑셀 표시형식 [숫자] (소수 자릿수 0, 음수 빨간 괄호)
NUM_FMT_INT = r"0_);[Red]\(0\)"
# 소수가 있는 값(예: 체중 81.1)은 소수점이 잘려 보이지 않도록 별도 형식
NUM_FMT_DEC = r"0.0###_);[Red]\(0.0###\)"

# ── 2단계: 열 추가 (앞 항목 다음에 새 열들을 순서대로 삽입) ──────────────
INSERT_COLUMNS = {
    "DEMO": [
        ("RPT_TY", ["보고구분"]),
        ("ADRSE_STUDY_TYP", ["연구구분"]),
        ("ADRSE_STUDY_LWPRT_TYP", ["기타시험구분"]),
        ("REPRT_CHANGE_CD", ["보고무효화"]),
        ("PRMRPT_TY", ["원보고자"]),
        ("PRMRPT_LWPRT_CD", ["기타의료전문가"]),
        ("SENDER_TY", ["보고자"]),
        ("SENDER_TY_MED_EXPERT", ["의료전문가"]),
        ("PTNT_OCCURSYMT_AGE_UNIT", ["연령단위"]),
        ("PTNT_AGRDE", ["연령대"]),
        ("PTNT_SEX", ["성별"]),
    ],
    "EVENT": [
        ("ADR_MEDDRA_ENG_NM", ["PT_한국어", "SOC_한국어"]),
        ("ADR_RESULT_CODE", ["최종관찰결과"]),
        ("WHOART_SEQ", ["중대한"]),
    ],
    "DRUG": [
        ("DRUG_GB", ["의약품구분"]),
        ("DRUG_CD", ["의약품명"]),
        ("DRUG_ACTION", ["의약품에 대한 조치"]),
    ],
}

# DRUG_EVENT 시트 끝에 추가할 열 (E열부터)
DRUG_EVENT_NEW_COLUMNS = [
    "재투여 및 재발현 여부", "부서접수번호", "보고자관리번호", "보고완료일", "의심약 구분",
    "의약품명", "신속보고(Y/0)", "부작용명", "LLT_한국어", "PT_한국어", "SOC_한국어",
    "반영여부", "발현일", "종료일", "최종관찰결과", "중대한 이상사례(Y/0)", "평가결과",
    "의약품에 대한 조치", "원보고자 구분", "기타 의료전문가 구분", "보고자 구분",
    "보고자 내 의료전문가 상세 구분", "환자 연령", "환자연령단위", "연령대", "환자생년",
    "환자성별", "보고구분", "연구구분", "기타연구구분", "문헌제목", "시험연구계획서번호",
    "최초추적그룹번호", "최초추적그룹 일련번호(seq)", "보고무효화/수정(삭제/수정/0)",
]

# ── 3단계: 코드 → 한글 ──────────────────────────────────────────────
CODE_MAPS = {
    "DEMO": {
        "보고구분": ("RPT_TY", {"1": "자발적 보고", "2": "시험/연구", "3": "기타", "4": "모름"}),
        "연구구분": ("ADRSE_STUDY_TYP", {"1": "임상시험", "2": "개별 환자 사용", "3": "기타시험"}),
        "기타시험구분": ("ADRSE_STUDY_LWPRT_TYP", {
            "1": "재심사보고 사용성적조사", "2": "재심사보고 시판 후 임상연구",
            "3": "재심사보고 특별조사", "4": "기타"}),
        "보고무효화": ("REPRT_CHANGE_CD", {"1": "삭제", "2": "수정"}),
        "원보고자": ("PRMRPT_TY", {
            "1": "의사, 치과의사, 한의사", "2": "약사, 한약사", "3": "기타 의료전문가",
            "4": "변호사", "5": "소비자 또는 기타 비의료전문가", "UNK": "모름"}),
        "기타의료전문가": ("PRMRPT_LWPRT_CD", {"1": "간호사", "2": "기타"}),
        "보고자": ("SENDER_TY", {
            "1": "제약회사", "2": "규제당국", "3": "의료전문가", "4": "지역의약품안전센터",
            "5": "WHO국제의약품모니터링센터", "6": "기타(예: 도매상 또는 관련 단체)",
            "7": "환자/소비자"}),
        "의료전문가": ("SENDER_TY_MED_EXPERT", {"1": "병의원", "2": "약국", "3": "보건소", "4": "기타"}),
        "연령단위": ("PTNT_OCCURSYMT_AGE_UNIT", {
            "105": "시간", "107": "일", "108": "주", "106": "월", "103": "년", "9": "연대"}),
        "연령대": ("PTNT_AGRDE", {
            "0": "태아", "1": "출생일~28일 미만", "2": "28일~24개월 미만",
            "3": "24개월~12세 미만", "4": "12세~19세 미만", "5": "19세~65세 미만",
            "6": "65세 이상"}),
        "성별": ("PTNT_SEX", {"1": "남성", "2": "여성"}),
    },
    "EVENT": {
        "최종관찰결과": ("ADR_RESULT_CODE", {
            "1": "회복됨", "2": "회복중", "3": "회복되지 않았음",
            "4": "회복했지만 후유증 있음", "5": "치명적", "0": "알려지지 않음"}),
    },
    "DRUG": {
        "의약품구분": ("DRUG_GB", {"1": "의심", "2": "병용", "3": "상호작용", "4": "비투여"}),
        "의약품에 대한 조치": ("DRUG_ACTION", {
            "1": "투여 중지", "2": "투여량 감소", "3": "투여량 증가",
            "4": "투여량 유지", "0": "모름", "9": "해당없음"}),
    },
    "DRUG_EVENT": {
        "재투여 및 재발현 여부": ("RE_DOSAGE_ALLERGY_YN", {
            "1": "재투여하고 이상반응이 다시 나타남", "2": "재투여했지만 이상반응 나타나지 않음",
            "3": "재투여했지만 결과는 알 수 없음", "4": "재투여하지 않아 해당사항 없음"}),
    },
    "ASSESSMENT": {
        "평가결과": ("EVALT_RESULT_CODE", {
            "1": "확실함", "2": "상당히 확실함", "3": "가능함",
            "4": "가능성 적음", "5": "평가 곤란", "6": "평가 불가"}),
    },
}

SERIOUS_COLUMNS = ["SE_DEATH", "SE_LIFE_MENACE", "SE_HSPTLZ_EXTN",
                   "SE_FNCT_DGRD", "SE_ANMLY", "SE_ETC_IMPRTNC_SITTN"]

# ── 4단계: DRUG_EVENT 정보 불러오기 ─────────────────────────────────
# (대상 열, 원본 시트, 원본 열 이름, 연결 키)   ※ 주석은 기존 엑셀 수식
#   키 "K"   : KAERS_NO
#   키 "KD"  : KAERS_NO + DRUG_SEQ
#   키 "KA"  : KAERS_NO + ADR_SEQ
#   키 "KDA" : KAERS_NO + DRUG_SEQ + ADR_SEQ
DRUG_EVENT_LOOKUPS = [
    ("부서접수번호", "DEMO", "DEPT_RECEIPT_NO", "K"),          # =INDEX(DEMO!B:B,MATCH(A2,DEMO!A:A,0))  ※괄호 오류 수정
    ("보고자관리번호", "DEMO", "SFRPNO", "K"),                  # =INDEX(DEMO!C:C,MATCH(A2,DEMO!A:A,0))  ※괄호 오류 수정
    ("보고완료일", "DEMO", "RPT_DL_DT", "K"),                   # =INDEX(DEMO!D:D,MATCH(A2,DEMO!A:A,0))  ※괄호 오류 수정
    ("의심약 구분", "DRUG", "의약품구분", "KD"),                # =INDEX(DRUG!D:D,MATCH(1,(DRUG!A:A=A2)*(DRUG!B:B=B2),0))
    ("의약품명", "DRUG", "의약품명", "KD"),                     # =INDEX(DRUG!F:F,MATCH(1,(DRUG!A:A=A2)*(DRUG!B:B=B2),0))
    ("신속보고(Y/0)", "DEMO", "QCK_RPT_YN", "K"),               # =INDEX(DEMO!P:P,MATCH(A2,DEMO!A:A,0))
    ("부작용명", "EVENT", "ADR_MEDDRA_KOR_NM", "KA"),           # =INDEX(EVENT!C:C,MATCH(1,(EVENT!A:A=A2)*(EVENT!B:B=C2),0))
    ("PT_한국어", "EVENT", "PT_한국어", "KA"),                  # =INDEX(EVENT!E:E, ...)
    ("SOC_한국어", "EVENT", "SOC_한국어", "KA"),                # =INDEX(EVENT!F:F, ...)
    ("발현일", "EVENT", "ADR_START_DT", "KA"),                  # =INDEX(EVENT!G:G, ...)
    ("종료일", "EVENT", "ADR_END_DT", "KA"),                    # =INDEX(EVENT!H:H, ...)
    ("최종관찰결과", "EVENT", "최종관찰결과", "KA"),            # =INDEX(EVENT!J:J, ...)
    ("중대한 이상사례(Y/0)", "EVENT", "중대한", "KA"),          # =INDEX(EVENT!N:N, ...)
    ("평가결과", "ASSESSMENT", "평가결과", "KDA"),              # =INDEX(ASSESSMENT!E:E,MATCH(1,(..A=A2)*(..B=B2)*(..C=C2),0))
    ("의약품에 대한 조치", "DRUG", "의약품에 대한 조치", "KD"), # =INDEX(DRUG!J:J, ...)
    ("원보고자 구분", "DEMO", "원보고자", "K"),                 # =INDEX(DEMO!U:U, ...)
    ("기타 의료전문가 구분", "DEMO", "기타의료전문가", "K"),    # =INDEX(DEMO!W:W, ...)
    ("보고자 구분", "DEMO", "보고자", "K"),                     # =INDEX(DEMO!Y:Y, ...)
    ("보고자 내 의료전문가 상세 구분", "DEMO", "의료전문가", "K"),  # =INDEX(DEMO!AA:AA, ...)
    ("환자 연령", "DEMO", "PTNT_OCCURSYMT_AGE", "K"),           # =INDEX(DEMO!AB:AB, ...)
    ("환자연령단위", "DEMO", "연령단위", "K"),                  # =INDEX(DEMO!AD:AD, ...)
    ("연령대", "DEMO", "연령대", "K"),                          # =INDEX(DEMO!AF:AF, ...)
    ("환자생년", "DEMO", "PTNT_BRTYR_YYYY", "K"),               # =INDEX(DEMO!AG:AG, ...)
    ("환자성별", "DEMO", "성별", "K"),                          # =INDEX(DEMO!AI:AI, ...)
    ("보고구분", "DEMO", "보고구분", "K"),                      # =INDEX(DEMO!G:G, ...)
    ("연구구분", "DEMO", "연구구분", "K"),                      # =INDEX(DEMO!I:I, ...)
    ("기타연구구분", "DEMO", "기타시험구분", "K"),              # =INDEX(DEMO!K:K, ...)
    ("문헌제목", "DEMO", "LTRTRE_INFO", "K"),                   # =INDEX(DEMO!L:L, ...)
    ("시험연구계획서번호", "DEMO", "CLIENT_STUDY_NO", "K"),     # =INDEX(DEMO!M:M, ...)
    ("최초추적그룹번호", "GROUP", "GROUP", "K"),                # =INDEX(GROUP!B:B,MATCH(A2,GROUP!A:A,0))
    ("최초추적그룹 일련번호(seq)", "GROUP", "SEQ", "K"),        # =INDEX(GROUP!C:C,MATCH(A2,GROUP!A:A,0))
    ("보고무효화/수정(삭제/수정/0)", "DEMO", "보고무효화", "K"),  # =INDEX(DEMO!S:S, ...)
]
# LLT_한국어, 반영여부 : 제공된 수식이 없어 공란

KEY_COLUMNS = {
    "K": ["KAERS_NO"],
    "KD": ["KAERS_NO", "DRUG_SEQ"],
    "KA": ["KAERS_NO", "ADR_SEQ"],
    "KDA": ["KAERS_NO", "DRUG_SEQ", "ADR_SEQ"],
}

NA = "#N/A"   # 엑셀에서 MATCH 가 실패했을 때와 같은 표시


# ════════════════════════════════════════════════════════════════════
# 데이터 구조: 시트 = {"headers": [...], "rows": [[...], ...]}
#              README 는 {"lines": [...]}
# ════════════════════════════════════════════════════════════════════
class BotError(Exception):
    pass


def log(msg: str) -> None:
    print(msg, flush=True)


def col_index(sheet: dict, name: str, sheet_name: str) -> int:
    try:
        return sheet["headers"].index(name)
    except ValueError:
        raise BotError(f"[{sheet_name}] 시트에서 '{name}' 열을 찾을 수 없습니다.")


def code_key(v) -> str | None:
    """코드값 비교용 정규화: 1, 1.0, '1', ' 1 ', '001' → '1'"""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    if s == "":
        return None
    if re.fullmatch(r"-?\d+\.0+", s):
        s = s.split(".")[0]
    if re.fullmatch(r"\d+", s):   # 앞자리 0 제거: '00103' → '103'
        s = str(int(s))
    return s.upper()


def match_key(v):
    """조회 키 비교용 정규화 (엑셀 MATCH 처럼 숫자 1 과 문자 '1' 을 같은 값으로 취급)"""
    k = code_key(v)
    return k if k is not None else ""


# ════════════════════════════════════════════════════════════════════
# 1단계: txt 읽기
# ════════════════════════════════════════════════════════════════════
NUM_RE = re.compile(r"-?\d+(\.\d+)?")


def is_numeric_text(s: str) -> bool:
    if not NUM_RE.fullmatch(s):
        return False
    digits = s.lstrip("-").split(".")[0]
    if len(digits) > 1 and digits.startswith("0"):   # 앞자리 0 이 있는 번호는 문자로 유지
        return False
    if len(digits) > 15:                              # 엑셀 숫자 정밀도(15자리) 초과는 문자로 유지
        return False
    return True


def to_number(s: str):
    return float(s) if "." in s else int(s)


def convert_types(headers: list[str], raw_rows: list[list[str]]) -> list[list]:
    """엑셀 파워쿼리처럼 열 단위로 형식을 정함
       - 값이 모두 숫자인 열 → 숫자, 빈 칸은 빈 셀(None)
       - 그 외 열 → 문자, 빈 칸은 '' """
    ncol = len(headers)
    numeric = []
    for c in range(ncol):
        vals = [r[c].strip() for r in raw_rows if c < len(r) and r[c].strip() != ""]
        numeric.append(bool(vals) and all(is_numeric_text(v) for v in vals))
    out = []
    for r in raw_rows:
        r = (r + [""] * ncol)[:ncol]
        row = []
        for c, v in enumerate(r):
            if numeric[c]:
                v2 = v.strip()
                row.append(to_number(v2) if v2 else None)
            else:
                row.append(v)
        out.append(row)
    return out


def detect_delimiter(text: str) -> str:
    first = text.splitlines()[0] if text else ""
    counts = {d: first.count(d) for d in ["\t", "|", ",", ";"]}
    best = max(counts, key=counts.get)
    return best if counts[best] > 0 else "\t"


def find_txt_files(folder: Path) -> dict[str, Path]:
    found: dict[str, Path] = {}
    names_by_len = sorted(SHEET_ORDER, key=len, reverse=True)
    for f in sorted(folder.glob("*.txt")):
        stem = f.stem.upper()
        for name in names_by_len:
            if re.search(rf"(^|[^A-Z0-9]){re.escape(name)}([^A-Z0-9]|$)", stem):
                if name in found:
                    raise BotError(f"[{name}] 에 해당하는 txt 가 2개 이상입니다: {found[name].name}, {f.name}")
                found[name] = f
                break
    missing = [n for n in SHEET_ORDER if n not in found]
    if missing:
        raise BotError("다음 txt 파일을 찾지 못했습니다: " + ", ".join(missing)
                       + f"\n  (폴더: {folder})")
    return found


def read_txt_folder(folder: Path) -> dict:
    files = find_txt_files(folder)
    data = {}
    for name in SHEET_ORDER:
        path = files[name]
        text = path.read_text(encoding="utf-8-sig")       # 65001: 유니코드(UTF-8)
        if name == "README":
            lines = text.splitlines()
            while lines and not lines[-1].strip():
                lines.pop()
            data[name] = {"lines": [ln if ln.strip() else None for ln in lines]}
        else:
            delim = detect_delimiter(text)
            reader = list(csv.reader(text.splitlines(), delimiter=delim))
            reader = [r for r in reader if any(x.strip() for x in r)]
            if not reader:
                raise BotError(f"{path.name} 파일이 비어 있습니다.")
            headers = [h.strip() for h in reader[0]]
            data[name] = {"headers": headers, "rows": convert_types(headers, reader[1:])}
        log(f"   - {path.name:<40} → [{name}] {len(data[name].get('rows', data[name].get('lines')))}행")
    return data


# ════════════════════════════════════════════════════════════════════
# 2단계: 열 추가
# ════════════════════════════════════════════════════════════════════
def insert_after(sheet: dict, sheet_name: str, after: str, new_cols: list[str]) -> None:
    if new_cols[0] in sheet["headers"]:
        return   # 이미 있음
    pos = col_index(sheet, after, sheet_name) + 1
    sheet["headers"][pos:pos] = new_cols
    for r in sheet["rows"]:
        r[pos:pos] = [None] * len(new_cols)


def append_cols(sheet: dict, new_cols: list[str]) -> None:
    for c in new_cols:
        if c not in sheet["headers"]:
            sheet["headers"].append(c)
            for r in sheet["rows"]:
                r.append(None)


def step2_add_columns(data: dict) -> None:
    for sheet_name, rules in INSERT_COLUMNS.items():
        for after, cols in rules:
            insert_after(data[sheet_name], sheet_name, after, cols)
    append_cols(data["DRUG_EVENT"], DRUG_EVENT_NEW_COLUMNS)
    append_cols(data["ASSESSMENT"], ["평가결과"])


# ════════════════════════════════════════════════════════════════════
# 3단계: 코드 → 한글
# ════════════════════════════════════════════════════════════════════
def readme_item_code(data: dict) -> str | None:
    for ln in data["README"]["lines"]:
        if ln and "품목코드" in ln:
            m = re.search(r"품목코드\s*:\s*([0-9A-Za-z]+)", ln)
            if m:
                return m.group(1)
    return None


def step3_fill_codes(data: dict, drug_display_name: str) -> None:
    unknown: dict[tuple, set] = {}
    for sheet_name, maps in CODE_MAPS.items():
        sh = data[sheet_name]
        for target, (source, mapping) in maps.items():
            ti = col_index(sh, target, sheet_name)
            si = col_index(sh, source, sheet_name)
            for r in sh["rows"]:
                k = code_key(r[si])
                if k is None:
                    r[ti] = None
                elif k in mapping:
                    r[ti] = mapping[k]
                else:
                    r[ti] = None
                    unknown.setdefault((sheet_name, source), set()).add(k)

    # EVENT [중대한]: SE_ 6개 열 중 하나라도 Y 이면 Y
    ev = data["EVENT"]
    ti = col_index(ev, "중대한", "EVENT")
    sidx = [col_index(ev, c, "EVENT") for c in SERIOUS_COLUMNS]
    for r in ev["rows"]:
        r[ti] = "Y" if any(str(r[i] or "").strip().upper() == "Y" for i in sidx) else None

    # DRUG [의약품명]: DRUG_CD 가 README 의 요청 품목코드와 같으면 제품명 입력
    code = readme_item_code(data)
    dr = data["DRUG"]
    ti = col_index(dr, "의약품명", "DRUG")
    ci = col_index(dr, "DRUG_CD", "DRUG")
    if code is None:
        log("   ! README 에서 '요청 품목코드'를 찾지 못해 DRUG [의약품명]을 비워 둡니다.")
    else:
        n = 0
        for r in dr["rows"]:
            if code_key(r[ci]) == code.upper():
                r[ti] = drug_display_name
                n += 1
        log(f"   - DRUG [의약품명]: 품목코드 {code} → '{drug_display_name}' ({n}행)")

    for (sh, col), vals in unknown.items():
        log(f"   ! [{sh}] {col} 에 변환표에 없는 코드가 있어 빈칸으로 두었습니다: {sorted(vals)}")


# ════════════════════════════════════════════════════════════════════
# 4단계: DRUG_EVENT 정보 불러오기 (파이썬에서 직접 계산)
# ════════════════════════════════════════════════════════════════════
def build_index(sheet: dict, sheet_name: str, key_type: str) -> dict:
    """키 → 첫 번째로 일치하는 행 (엑셀 MATCH 와 동일하게 첫 행 사용)"""
    kidx = [col_index(sheet, c, sheet_name) for c in KEY_COLUMNS[key_type]]
    idx = {}
    for r in sheet["rows"]:
        k = tuple(match_key(r[i]) for i in kidx)
        idx.setdefault(k, r)
    return idx


def excel_value(v):
    """엑셀 INDEX 가 빈 셀을 가리키면 0 을 돌려주는 동작을 그대로 재현"""
    return 0 if v is None else v


def step4_lookup(data: dict) -> None:
    de = data["DRUG_EVENT"]
    de_key_idx = {t: [col_index(de, c, "DRUG_EVENT") for c in cols] for t, cols in KEY_COLUMNS.items()}
    cache = {}
    plan = []
    for target, src_sheet, src_col, key_type in DRUG_EVENT_LOOKUPS:
        ck = (src_sheet, key_type)
        if ck not in cache:
            cache[ck] = build_index(data[src_sheet], src_sheet, key_type)
        plan.append((col_index(de, target, "DRUG_EVENT"),
                     col_index(data[src_sheet], src_col, src_sheet),
                     cache[ck], key_type, target))
    na_count: dict[str, int] = {}
    for r in de["rows"]:
        for ti, si, index, key_type, target in plan:
            k = tuple(match_key(r[i]) for i in de_key_idx[key_type])
            src = index.get(k)
            if src is None:
                r[ti] = NA
                na_count[target] = na_count.get(target, 0) + 1
            else:
                r[ti] = excel_value(src[si])
    for t, n in na_count.items():
        log(f"   ! DRUG_EVENT [{t}]: 일치하는 행을 찾지 못한 {n}건은 #N/A 로 표시")


# ════════════════════════════════════════════════════════════════════
# 엑셀 쓰기 / 읽기
# ════════════════════════════════════════════════════════════════════
def apply_number_format(cell) -> None:
    v = cell.value
    if isinstance(v, float) and not v.is_integer():
        cell.number_format = NUM_FMT_DEC
    else:
        cell.number_format = NUM_FMT_INT


def write_rows(ws, headers, rows, number_format=True) -> None:
    ws.append(headers)
    for r in rows:
        ws.append(r)
    if number_format:
        for row in ws.iter_rows():
            for cell in row:
                apply_number_format(cell)


def add_table(ws, name: str, ncols: int, nrows: int) -> None:
    if nrows < 1 or ncols < 1:
        return
    ref = f"A1:{get_column_letter(ncols)}{nrows + 1}"
    t = Table(displayName=re.sub(r"[^A-Za-z0-9_]", "_", name), ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)


def save_full_workbook(data: dict, path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name in SHEET_ORDER:
        ws = wb.create_sheet(name)
        sh = data[name]
        if name == "README":
            for ln in sh["lines"]:
                ws.append([ln])
            continue
        write_rows(ws, sh["headers"], sh["rows"])
        add_table(ws, name, len(sh["headers"]), len(sh["rows"]))
    wb.save(path)
    log(f"   ✔ 저장: {path.name}")


def load_full_workbook(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    data = {}
    for name in SHEET_ORDER:
        if name not in wb.sheetnames:
            raise BotError(f"{path.name} 에 [{name}] 시트가 없습니다.")
        ws = wb[name]
        if name == "README":
            lines = [r[0] for r in ws.iter_rows(values_only=True)]
            data[name] = {"lines": lines}
            continue
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        while rows and all(v is None for v in rows[-1]):
            rows.pop()
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        while headers and headers[-1] == "":
            headers.pop()
        n = len(headers)
        body = [(r + [None] * n)[:n] for r in rows[1:]]
        # 엑셀 파일에는 빈 문자열('')이 저장되지 않으므로, 원본 문자형 열의 빈칸을 '' 로 되돌림
        # (엑셀 INDEX 결과: 빈 문자열 → 빈칸, 빈 셀 → 0  을 그대로 재현하기 위함)
        added = {c for _, cols in INSERT_COLUMNS.get(name, []) for c in cols}
        if name == "DRUG_EVENT":
            added |= set(DRUG_EVENT_NEW_COLUMNS)
        if name == "ASSESSMENT":
            added.add("평가결과")
        for c, h in enumerate(headers):
            if h in added:
                continue
            if not any(isinstance(r[c], (int, float)) for r in body):
                for r in body:
                    if r[c] is None:
                        r[c] = ""
        data[name] = {"headers": headers, "rows": body}
    return data


def safe_sheet_name(s: str) -> str:
    s = re.sub(r"[\[\]:*?/\\]", "_", s)
    return s[:31]


def save_values_file(data: dict, path: Path) -> None:
    """5단계: DRUG_EVENT 를 값으로 새 파일(Sheet1)에, 숫자 형식, D2 틀 고정"""
    de = data["DRUG_EVENT"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    write_rows(ws, de["headers"], de["rows"])
    ws.freeze_panes = "D2"
    wb.save(path)
    log(f"   ✔ 저장: {path.name}")


def save_split_file(data: dict, path: Path, sheet_prefix: str, drug_display_name: str) -> dict:
    """6단계: 원본 / 제품명_의심 / 제품명_병용"""
    de = data["DRUG_EVENT"]
    h = de["headers"]
    gi = col_index(de, "의심약 구분", "DRUG_EVENT")
    ni = col_index(de, "의약품명", "DRUG_EVENT")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "원본"
    write_rows(ws, h, de["rows"])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(h))}{len(de['rows']) + 1}"
    ws.freeze_panes = "D2"

    counts = {}
    for kind in ["의심", "병용"]:
        rows = [r for r in de["rows"] if r[gi] == kind and r[ni] == drug_display_name]
        s = wb.create_sheet(safe_sheet_name(f"{sheet_prefix}_{kind}"))
        write_rows(s, h, rows)
        s.freeze_panes = "D2"
        counts[kind] = len(rows)
    wb.save(path)
    log(f"   ✔ 저장: {path.name}  (의심 {counts['의심']}행, 병용 {counts['병용']}행)")
    return counts


# ════════════════════════════════════════════════════════════════════
# 실행 흐름
# ════════════════════════════════════════════════════════════════════
def out_name(product: str, no: str, date: str) -> str:
    product = re.sub(r'[\\/:*?"<>|]', "_", product)   # 파일 이름에 쓸 수 없는 문자 (예: 2/500 → 2_500)
    return f"원시자료_{product}_{no}_{date}.xlsx"


def run_steps_1_3(txt_dir: Path, out_dir: Path, product: str, drug_name: str, date: str) -> dict:
    log("\n[1단계] txt 불러오기")
    data = read_txt_folder(txt_dir)
    save_full_workbook(data, out_dir / out_name(product, "01", date))

    log("\n[2단계] 열 추가")
    step2_add_columns(data)
    save_full_workbook(data, out_dir / out_name(product, "02", date))

    log("\n[3단계] 코드 → 한글 정보 입력")
    step3_fill_codes(data, drug_name)
    save_full_workbook(data, out_dir / out_name(product, "03", date))
    return data


def run_steps_4_6(data: dict, out_dir: Path, product: str, drug_name: str,
                  sheet_prefix: str, date: str) -> None:
    ev = data["EVENT"]
    pi = col_index(ev, "PT_한국어", "EVENT")
    blank_pt = sum(1 for r in ev["rows"] if r[pi] in (None, ""))
    if blank_pt:
        log(f"\n   ※ EVENT 시트의 PT_한국어가 비어 있는 행이 {blank_pt}개 있습니다. (DRUG_EVENT 에 0 으로 표시)")

    log("\n[4단계] DRUG_EVENT 정보 불러오기")
    step4_lookup(data)
    save_full_workbook(data, out_dir / out_name(product, "10", date))

    log("\n[5단계] 값으로 새 파일 만들기")
    save_values_file(data, out_dir / out_name(product, "20", date))

    log("\n[6단계] 의심 / 병용 시트 나누기")
    save_split_file(data, out_dir / out_name(product, "30", date), sheet_prefix, drug_name)


# ── 화면 입력 ─────────────────────────────────────────────────────
def ask(prompt: str, default: str | None = None) -> str:
    tip = f" [{default}]" if default else ""
    while True:
        v = input(f"{prompt}{tip}: ").strip().strip('"')
        if v:
            return v
        if default is not None:
            return default


def main() -> None:
    print("=" * 60)
    print(f" 원시자료 정리봇  ({VERSION})")
    print("=" * 60)
    print(" [1] 전체 실행     (txt → _01 ~ _30)")
    print(" [2] 1~3단계만     (txt → _01, _02, _03)  ※ 이후 PT/SOC 직접 입력")
    print(" [3] 4~6단계만     (PT/SOC 입력한 _03 파일 → _10, _20, _30)")
    mode = ask("작업 번호", "1")
    if mode not in {"1", "2", "3"}:
        raise BotError("작업 번호는 1, 2, 3 중 하나를 입력하세요.")

    if mode in {"1", "2"}:
        src = Path(ask("txt 파일 13개가 들어 있는 폴더 경로"))
        if not src.is_dir():
            raise BotError(f"폴더가 없습니다: {src}")
    else:
        src = Path(ask("PT/SOC 입력을 마친 _03 엑셀 파일 경로"))
        if not src.is_file():
            raise BotError(f"파일이 없습니다: {src}")

    loaded = None
    product = ask("제품명 (파일 이름용, 예: 그리타존정15mg)")
    if mode == "3":
        log("\n[_03 파일 읽기]")
        loaded = load_full_workbook(src)
        dr = loaded["DRUG"]
        ni = col_index(dr, "의약품명", "DRUG")
        names = sorted({str(r[ni]) for r in dr["rows"] if r[ni] not in (None, "", 0)})
        if len(names) == 1:
            drug_name = names[0]
            log(f"   - DRUG [의약품명]: '{drug_name}'")
        else:
            drug_name = ask("DRUG 시트 [의약품명] 중 의심/병용으로 나눌 제품 이름", names[0] if names else None)
    else:
        drug_name = ask("DRUG 시트 [의약품명]에 표시할 이름 (예: 그리타존정15밀리그램)", product)
    sheet_prefix = ask("_30 파일 시트 이름 앞부분 (예: 그리타존정)", product) if mode != "2" else product
    date = ask("파일 날짜 YYYYMMDD", dt.date.today().strftime("%Y%m%d"))
    if not re.fullmatch(r"\d{8}", date):
        raise BotError("날짜는 YYYYMMDD 8자리 숫자로 입력하세요.")
    out_dir = Path(ask("결과 저장 폴더", str(src if src.is_dir() else src.parent)))
    out_dir.mkdir(parents=True, exist_ok=True)

    if mode == "1":
        data = run_steps_1_3(src, out_dir, product, drug_name, date)
        run_steps_4_6(data, out_dir, product, drug_name, sheet_prefix, date)
    elif mode == "2":
        run_steps_1_3(src, out_dir, product, drug_name, date)
        log("\n다음 작업: _03 파일 EVENT 시트에 PT_한국어 / SOC_한국어 입력 후 [3] 번으로 실행하세요.")
    else:
        run_steps_4_6(loaded, out_dir, product, drug_name, sheet_prefix, date)

    log(f"\n완료되었습니다.  저장 위치: {out_dir.resolve()}")
    if sys.stdin.isatty() and hasattr(os, "startfile"):
        try:
            os.startfile(out_dir.resolve())   # 결과 폴더 열기
        except OSError:
            pass


if __name__ == "__main__":
    try:
        main()
    except BotError as e:
        print(f"\n[오류] {e}")
    except PermissionError as e:
        print(f"\n[오류] 파일을 저장할 수 없습니다. 같은 이름의 엑셀 파일이 열려 있으면 닫고 다시 실행하세요.\n  {e}")
    except (KeyboardInterrupt, EOFError):
        print("\n중단되었습니다.")
    except Exception:
        # 예상하지 못한 오류도 창이 바로 닫히지 않도록 화면에 표시
        print("\n[오류] 예상하지 못한 오류가 발생했습니다. 아래 내용을 개발자에게 전달해 주세요.\n")
        traceback.print_exc()
    try:
        input("\nEnter 를 누르면 종료합니다.")
    except EOFError:
        pass
